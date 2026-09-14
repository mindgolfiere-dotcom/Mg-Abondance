#!/usr/bin/env python3
"""Reconstruit les 2 slides de synthese Formations 2026 dans le template EssilorLuxottica."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_MARK

SRC = "formation.pptx"
OUT = "Synthese_Formations_2026.pptx"

# --- Charte EssilorLuxottica extraite du deck source -------------------------
NAVY = RGBColor(0x1E, 0x27, 0x61)
ORANGE = RGBColor(0xF2, 0x99, 0x4A)
GREY = RGBColor(0x66, 0x70, 0x85)
STEEL = RGBColor(0x5C, 0x8D, 0xB8)   # accent1 du theme
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
CARD = RGBColor(0xF4, 0xF6, 0xFA)
GRID = RGBColor(0xE4, 0xE8, 0xF0)

SERIF, SANS = "Cambria", "Calibri"

# --- Donnees YTD consolidees depuis le MBR (Jan -> Aout 2026) ----------------
MOIS = ["Jan", "Fév", "Mar", "Avr", "Mai", "Juin", "Juil-Août"]
PARTENAIRES = [19, 14, 19, 18, 13, 16, 38]   # total 137
INTERNES = [2, 2, 2, 4, 8, 2, 4]             # total  24
OPHTAS = [4, 4, 6, 13, 2, 4, 18]             # total  51

# --- Plan Q4 2026 -----------------------------------------------------------
Q4_MOIS = ["SEPT", "OCT", "NOV", "DÉC"]
Q4_INTERNES = [1, 1, 2, 1]                   # total 5
Q4_WEBINAIRES = [1, 1, 2, 1]                 # total 5
THEMES = [("PAL Essilor", 1), ("Nikon Z FSV", 1), ("Varilux Shift", 2),
          ("Kodak FSV", 2), ("Stellest", 4)]  # ordre ascendant = haut en bas


# --- Helpers ----------------------------------------------------------------
def clear_slide(slide, keep):
    """Vide la slide en conservant les elements de pied de page du template."""
    tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        if shape.name in keep:
            continue
        tree.remove(shape._element)
    part = slide.part
    for rId, rel in list(part.rels.items()):
        if rel.reltype.rsplit("/", 1)[-1] in ("chart", "oleObject", "image"):
            part.drop_rel(rId)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def run(para, text, size, *, bold=False, font=SANS, color=GREY):
    r = para.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = color
    return r


def card(slide, x, y, w, h, fill=CARD):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def dot(slide, x, y, d=0.13, color=ORANGE):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def header(slide, eyebrow, title, subtitle):
    tf = textbox(slide, 0.21, 0.25, 12.7, 0.95)
    p = tf.paragraphs[0]
    run(p, eyebrow, 23, bold=True, font="Arial", color=BLACK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(4)
    run(p2, title, 18, bold=True, font=SERIF, color=NAVY)
    run(p2, subtitle, 18, font=SANS, color=GREY)


def kpi_row(slide, items, y=1.32, h=1.22, x0=0.5, total_w=12.33, gap=0.23):
    n = len(items)
    w = (total_w - gap * (n - 1)) / n
    for i, (value, caption) in enumerate(items):
        x = x0 + i * (w + gap)
        card(slide, x, y, w, h)
        dot(slide, x + w - 0.30, y + 0.17)
        tf = textbox(slide, x + 0.18, y + 0.13, w - 0.55, 0.56)
        run(tf.paragraphs[0], value, 34, bold=True, font=SERIF, color=NAVY)
        tf = textbox(slide, x + 0.18, y + 0.70, w - 0.34, h - 0.82)
        run(tf.paragraphs[0], caption, 10.5, font=SANS, color=GREY)


def style_chart(chart, *, legend=True):
    chart.has_title = False          # le titre est un textbox, pas le nom de serie
    chart.font.name = SANS
    chart.font.size = Pt(10)
    chart.font.color.rgb = GREY
    if legend:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(9.5)
        chart.legend.font.color.rgb = GREY
    else:
        chart.has_legend = False


def chart_title(slide, x, y, w, text):
    tf = textbox(slide, x, y, w, 0.3)
    run(tf.paragraphs[0], text, 13, bold=True, font=SERIF, color=NAVY)


def quiet_axes(chart, *, value_max=None, value_unit=None, hide_value=False):
    va = chart.value_axis
    va.has_major_gridlines = not hide_value
    if not hide_value:
        va.major_gridlines.format.line.color.rgb = GRID
        va.major_gridlines.format.line.width = Pt(0.75)
    va.has_minor_gridlines = False
    va.visible = not hide_value
    va.format.line.fill.background()
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minor_tick_mark = XL_TICK_MARK.NONE
    va.minimum_scale = 0
    if value_max is not None:
        va.maximum_scale = value_max
    if value_unit is not None:
        va.major_unit = value_unit
    va.tick_labels.font.size = Pt(9)
    va.tick_labels.font.color.rgb = GREY

    ca = chart.category_axis
    ca.has_major_gridlines = False
    ca.has_minor_gridlines = False
    ca.format.line.color.rgb = GRID
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(10)
    ca.tick_labels.font.bold = True
    ca.tick_labels.font.color.rgb = NAVY


def series_color(series, color):
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = color
    series.format.line.fill.background()


def label_font(series, color, size=10, bold=True):
    dl = series.data_labels
    dl.show_value = True
    dl.font.size = Pt(size)
    dl.font.bold = bold
    dl.font.color.rgb = color
    dl.font.name = SANS
    return dl


def navy_strip(slide, y, h, blocks, title):
    """Bandeau navy: un libelle + des paires (chiffre, legende)."""
    card(slide, 0.5, y, 12.33, h, fill=NAVY)
    tf = textbox(slide, 0.78, y, 2.45, h, anchor=MSO_ANCHOR.MIDDLE)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    run(tf.paragraphs[0], title, 14, bold=True, font=SERIF, color=WHITE)
    x = 3.30
    slot = (12.83 - x - 0.2) / len(blocks)
    for value, caption in blocks:
        tf = textbox(slide, x, y, slot - 0.15, h, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run(p, value + "  ", 15, bold=True, font=SERIF, color=ORANGE)
        run(p, caption, 11, font=SANS, color=WHITE)
        x += slot


# ============================================================================
prs = Presentation(SRC)
KEEP = {"Espace réservé du pied de page 2", "Google Shape;2474;p12"}

# ---------------------------------------------------------------- SLIDE 1 ---
s1 = prs.slides[0]
clear_slide(s1, KEEP)

header(s1, "EDUCATION TO GO | MOROCCO",
       "Bilan Formations YTD", " — Janvier à Août 2026")

kpi_row(s1, [
    ("212", "Actions de formation réalisées sur 8 mois"),
    ("137", "Formations partenaires en magasin, dont 75 Experts"),
    ("51", "Visites ophtalmologues, + 183 rencontrés au Congrès SMO"),
    ("8", "Villes couvertes · 6 Roadshows Kodak · 2 webinaires"),
])

# Graphique 1 : colonnes groupees par mois (chaque valeur reste lisible)
chart_title(s1, 0.72, 2.78, 7.0, "Activité de formation par mois")
cd = CategoryChartData()
cd.categories = MOIS
cd.add_series("Formations partenaires", PARTENAIRES)
cd.add_series("Visites ophtalmologues", OPHTAS)
cd.add_series("Formations internes", INTERNES)
gf = s1.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(3.12),
                         Inches(7.55), Inches(2.95), cd)
c1 = gf.chart
style_chart(c1)
plot = c1.plots[0]
plot.gap_width = 60
plot.overlap = -10
for srs, col in ((c1.series[0], NAVY), (c1.series[1], ORANGE),
                 (c1.series[2], STEEL)):
    series_color(srs, col)
    label_font(srs, col, size=9).position = XL_LABEL_POSITION.OUTSIDE_END
quiet_axes(c1, value_max=45, value_unit=15)

# Graphique 2 : donut de repartition
chart_title(s1, 8.5, 2.78, 4.3, "Répartition des 212 actions")
cd2 = CategoryChartData()
cd2.categories = ["Partenaires", "Ophtalmologues", "Internes"]
cd2.add_series("Actions", (sum(PARTENAIRES), sum(OPHTAS), sum(INTERNES)))
gf2 = s1.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(8.28), Inches(3.12),
                          Inches(4.55), Inches(2.95), cd2)
c2 = gf2.chart
style_chart(c2)
donut = c2.plots[0]
for pt, col in zip(donut.series[0].points, (NAVY, ORANGE, STEEL)):
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = col
    pt.format.line.color.rgb = WHITE
    pt.format.line.width = Pt(1.5)
dl = donut.data_labels
dl.show_value = True
dl.show_percentage = False
dl.font.size = Pt(11)
dl.font.bold = True
dl.font.name = SANS
dl.font.color.rgb = WHITE

navy_strip(s1, 6.20, 0.62, [
    ("120", "comptes ECP actifs"),
    ("90 %", "comptes internes actifs"),
    ("3", "séances de formation menées"),
], "Leonardo")

# ---------------------------------------------------------------- SLIDE 2 ---
s2 = prs.slides[1]
clear_slide(s2, KEEP)

header(s2, "EDUCATION TO GO | MOROCCO",
       "Plan de Formation Q4 2026", " — Septembre à Décembre")

kpi_row(s2, [
    ("10", "Sessions planifiées : 5 internes + 5 webinaires"),
    ("5", "Thèmes produits couverts, priorité Stellest"),
    ("7", "Groupes partenaires ciblés — Essilor Expert / Nikon"),
    ("100 %", "Formations obligatoires à compléter avant le 31/12"),
])

# Graphique 1 : sessions planifiees par mois
chart_title(s2, 0.72, 2.78, 5.6, "Sessions planifiées par mois")
cd3 = CategoryChartData()
cd3.categories = Q4_MOIS
cd3.add_series("Formations internes", Q4_INTERNES)
cd3.add_series("Webinaires", Q4_WEBINAIRES)
gf3 = s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.5), Inches(3.10),
                          Inches(6.0), Inches(2.42), cd3)
c3 = gf3.chart
style_chart(c3)
plot3 = c3.plots[0]
plot3.gap_width = 70
plot3.overlap = 100
for srs, col, lab in ((c3.series[0], NAVY, WHITE), (c3.series[1], ORANGE, NAVY)):
    series_color(srs, col)
    label_font(srs, lab, size=11).position = XL_LABEL_POSITION.CENTER
quiet_axes(c3, value_max=4, value_unit=1, hide_value=True)

tf = textbox(s2, 0.72, 5.60, 6.0, 0.42)
p = tf.paragraphs[0]
run(p, "Sept ", 9.5, bold=True, font=SANS, color=NAVY)
run(p, "Kodak FSV  ·  ", 9.5, font=SANS, color=GREY)
run(p, "Oct ", 9.5, bold=True, font=SANS, color=NAVY)
run(p, "Stellest  ·  ", 9.5, font=SANS, color=GREY)
run(p, "Nov ", 9.5, bold=True, font=SANS, color=NAVY)
run(p, "Varilux Shift, Nikon Z, PAL Essilor  ·  ", 9.5, font=SANS, color=GREY)
run(p, "Déc ", 9.5, bold=True, font=SANS, color=NAVY)
run(p, "Stellest", 9.5, font=SANS, color=GREY)

# Graphique 2 : barres horizontales par theme
chart_title(s2, 7.05, 2.78, 5.6, "Sessions par thème produit")
cd4 = CategoryChartData()
cd4.categories = [t for t, _ in THEMES]
cd4.add_series("Sessions", [v for _, v in THEMES])
gf4 = s2.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(6.83), Inches(3.10),
                          Inches(6.0), Inches(2.42), cd4)
c4 = gf4.chart
style_chart(c4, legend=False)
plot4 = c4.plots[0]
plot4.gap_width = 60
srs = c4.series[0]
series_color(srs, NAVY)
for i, pt in enumerate(srs.points):
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = ORANGE if THEMES[i][0] == "Stellest" else NAVY
    pt.format.line.fill.background()
label_font(srs, NAVY, size=10.5).position = XL_LABEL_POSITION.OUTSIDE_END
quiet_axes(c4, value_max=5, value_unit=1, hide_value=True)
c4.category_axis.tick_labels.font.size = Pt(10.5)

navy_strip(s2, 6.18, 0.64, [
    ("1 500 MAD", "1er prix"),
    ("1 000 MAD", "2e prix"),
    ("700 MAD", "3e prix"),
    ("100 %", "comptes actifs"),
], "Challenge Leonardo")

s1.notes_slide.notes_text_frame.text = (
    "Chiffres consolidés mois par mois depuis le MBR (Jan → Août 2026).\n"
    "Formations partenaires : 19+14+19+18+13+16+38 = 137.\n"
    "Visites ophtalmologues : 4+4+6+13+2+4+18 = 51.\n"
    "Formations internes : 2+2+2+4+8+2+4 = 24.  Total = 212 actions.\n"
    "Congrès SMO (183+), 8 villes, 6 Roadshows Kodak, 2 webinaires "
    "(Leonardo 34 participants, Confort visuel 72) et les indicateurs "
    "Leonardo proviennent du bilan consolidé DOP."
)
s2.notes_slide.notes_text_frame.text = (
    "5 formations internes : Kodak FSV (sept), Stellest (oct), Varilux Shift "
    "et Nikon Z FSV (nov), Stellest (déc).\n"
    "5 webinaires : Kodak FSV (sept), Stellest S1 (oct), PAL Essilor Intro et "
    "Varilux Shift (nov), Stellest S2 (déc).\n"
    "7 groupes partenaires ciblés : Lynx, Krys, Galerie, Comptoir, "
    "Santé Vision, Oudghiri, Laalej.\n"
    "Challenge Leonardo : le département le plus engagé remporte également un "
    "petit-déjeuner d'équipe. Session de rappel interne fin septembre."
)

prs.save(OUT)
print("OK ->", OUT)
